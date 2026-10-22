"""Turns the transcribed Markdown subset into a Word document."""

import io
import re

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

_INLINE = re.compile(r"(\*\*.+?\*\*|__.+?__|~~.+?~~)")
_HEADING = re.compile(r"^(#{1,3})\s+(.*)$")
_BULLET = re.compile(r"^[-*•]\s+(.*)$")
_NUMBERED = re.compile(r"^\d+[.)]\s+(.*)$")


def _set_language(doc: Document, lang: str = "lt-LT"):
    """Mark the text as Lithuanian so Word's spell checker uses the right dictionary."""
    rpr = doc.styles["Normal"].element.get_or_add_rPr()
    el = rpr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rpr.append(el)
    el.set(qn("w:val"), lang)


def _add_runs(paragraph, text: str):
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("__") and part.endswith("__") and len(part) > 4:
            paragraph.add_run(part[2:-2]).underline = True
        elif part.startswith("~~") and part.endswith("~~") and len(part) > 4:
            paragraph.add_run(part[2:-2]).font.strike = True
        else:
            paragraph.add_run(part)


def _add_page_text(doc: Document, text: str):
    paragraph_lines: list[str] = []

    def flush():
        if paragraph_lines:
            _add_runs(doc.add_paragraph(), " ".join(paragraph_lines))
            paragraph_lines.clear()

    for line in text.splitlines():
        line = line.strip()
        if not line:
            flush()
            continue
        if m := _HEADING.match(line):
            flush()
            heading = doc.add_heading(level=len(m.group(1)))
            _add_runs(heading, m.group(2))
        elif m := _BULLET.match(line):
            flush()
            _add_runs(doc.add_paragraph(style="List Bullet"), m.group(1))
        elif m := _NUMBERED.match(line):
            flush()
            _add_runs(doc.add_paragraph(style="List Number"), m.group(1))
        else:
            paragraph_lines.append(line)
    flush()


def build_docx(texts: list[str], title: str = "", page_breaks: bool = False) -> bytes:
    doc = Document()
    _set_language(doc)
    doc.styles["Normal"].font.size = Pt(12)
    if title.strip():
        doc.add_heading(title.strip(), level=0)
    for i, text in enumerate(texts):
        if i > 0 and page_breaks:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        _add_page_text(doc, text)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
